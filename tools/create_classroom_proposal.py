"""Create the illustrated course DOCX. Run create_course_diagrams.py first.
Usage: python create_classroom_proposal.py /path/to/output.docx
Dependencies: python-docx. Illustration generation also needs cairosvg.
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1];ART=ROOT/'docs/course-art'
d=Document();sec=d.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.62);sec.left_margin=sec.right_margin=Inches(.72)
normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',27),('Subtitle',13),('Heading 1',22),('Heading 2',14)]:
 st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(7);st.paragraph_format.space_before=Pt(10 if name=='Heading 2' else 0)
 st.paragraph_format.keep_with_next=True

def p(text,style=None):return d.add_paragraph(text,style)
def code(text):
 para=p(text);para.paragraph_format.space_after=Pt(10)
 for run in para.runs:run.font.name='DejaVu Sans Mono';run.font.size=Pt(10)
def h(text):d.add_heading(text,2)
def page(title):d.add_page_break();d.add_heading(title,1)
def boldline(label,text):
 para=d.add_paragraph();para.add_run(label+' ').bold=True;para.add_run(text);return para
def lablink(key,label):
 para=d.add_paragraph()
 link=OxmlElement('w:hyperlink')
 link.set(qn('r:id'),d.part.relate_to('http://localhost:8090/interactive/SharpHDL_WASM_Lab.html#'+key,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
 run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'075963');props.append(color);run.append(props)
 text=OxmlElement('w:t');text.text='Browser experiment: '+label;run.append(text);link.append(run);para._p.append(link)

def picture(name,caption,width=7):
 para=d.add_paragraph();para.paragraph_format.space_after=Pt(2);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 shape=para.add_run().add_picture(str(ART/(name+'.png')),width=Inches(width))
 shape._inline.docPr.set('descr',caption);shape._inline.docPr.set('title',name.replace('_',' '))
 para.paragraph_format.keep_with_next=True
 cap=p(caption);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.space_after=Pt(8)
 for run in cap.runs:run.font.size=Pt(10);run.italic=True

def table(rows,widths,size=10.5):
 t=d.add_table(rows=0,cols=len(widths));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for ri,row in enumerate(rows):
  cells=t.add_row().cells
  for ci,(cell,text) in enumerate(zip(cells,row)):
   cell.width=Inches(widths[ci]);cell.text=str(text);cell.vertical_alignment=1
   pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17365D' if ri==0 else 'EDF5F7' if ri%2==0 else 'FFFFFF');pr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for edge in ('top','left','bottom','right'):
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders);margins=OxmlElement('w:tcMar')
   for edge in ('top','left','bottom','right'):
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.05
    if ci==0 and len(str(text))<16:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:
     run.font.size=Pt(size)
     if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
  # Keep a row together and repeat headers when a long table spans a page.
  trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
 spacer=p('');spacer.paragraph_format.space_after=Pt(4);spacer.paragraph_format.line_spacing=Pt(1)
 for run in spacer.runs:run.font.size=Pt(1)
 return t

def mission(text):boldline('Your mission',text)
def tryit(text):h('Try it');p(text)
def ask(text):boldline('Tell your partner',text)
def note(text):
 para=boldline('Guide note',text)
 for run in para.runs:run.font.size=Pt(10)

# 1 Cover
p('Digital Hardware with SharpHDL','Title')
p('An illustrated workbook for curious builders around age ten','Subtitle')
p('Make numbers glow. Build a circuit that adds. Teach a register to remember. Then turn a clock into a blinking counter.')
picture('cover','A picture of discovery and teamwork. The pictured board is imaginative cover art, not the real board layout.')
h('How to use this workbook')
p('Work with a partner and a grown-up guide. Predict the answer before you press a button or change the code. Test your idea, then explain what you saw. Swap the builder and predictor jobs halfway through.')
table([['First explore','Then build','Finally explain'],['Bits and lights','Adders and memory','Counters and an assembly line'],['Number cards and gate puzzles','A small edit to real C# code','Your own useful circuit']],[2.3,2.3,2.4],11)
boldline('Name','________________________________   Partner ________________________________')
p('Eight one-hour lessons, plus printable number cards, a display key, a teacher guide, and answers. The first six lessons can be a complete shorter course. Language comparison is optional enrichment.')

# 2 Board and controls
page('Meet your little hardware workshop')
mission('Find the buttons, lights and number displays before writing any code.')
picture('board','A simplified Go Board feature map. It shows what each part does, not its exact physical position.')
table([['Button','What the supplied starter does','Try this'],['1','Clear B and the counter','Return the number to zero'],['2','Increment B once per press','Watch the binary LEDs change'],['3','Toggle the counter between run and pause','Start and stop counting'],['4','Toggle arithmetic or counter displays','Compare the two views']],[.65,3.75,2.6])
p('In arithmetic mode, the left digit shows B. The right digit shows B + 2, keeping four bits. The LEDs show B in binary. In counter mode, the two digits show an eight-bit count and the LEDs show its low four bits.')
tryit('Start with B = 0. Press button 2 three times. Predict the left digit, right digit and LEDs before looking. Hold the button down: it should count once, not over and over.')
note('Recommended pilot board: Nandland Go Board, iCE40 HX1K, four pushbuttons, four LEDs and two seven-segment digits. Maker list price checked 13 October 2026: $70. The board has pushbuttons rather than slide switches. Budget about $90 per pair with estimated cable/case allowance; tax, shipping and computers are extra. Validate one real board first. Sources are in the teacher guide.')

# 3 Lesson 1
page('Lesson one  Numbers made of lights')
mission('Build and read a number using four bits.')
p('A bit has two choices: 0 or 1. For our lights, 0 means OFF and 1 means ON. Each place has a value. Read the places from left to right: 8, 4, 2, 1.')
picture('binary_cards','ON contributes the card value. OFF contributes zero. The pattern 1101 represents thirteen.')
table([['Number','8 card','4 card','2 card','1 card','Binary'],['5',0,1,0,1,'0101'],['9',1,0,0,1,'1001'],['13',1,1,0,1,'1101'],['Your turn: 6','___','___','___','___','____']],[1.65,.88,.88,.88,.88,1.83])
tryit('Use the printable cards near the end of the workbook. Turn the cards ON or OFF to make 3, 6 and 12. Enter the same numbers on the board with button 2 and compare the four LEDs.')
ask('Why does 1000 mean eight rather than one thousand? How many patterns can four bits make?')
boldline('I can explain','A bit has ______ choices. Four bits can show ______ different patterns.')
note('Success: the learner correctly makes and reads three numbers from 0 to 15. Sixteen combinations include zero. No code is needed for this lesson.')

# 4 Lesson 2
page('Lesson two  Gates make decisions')
mission('Predict whether a light will be ON using an AND or OR rule.')
p('Imagine two friends holding input cards. AND says both must show 1. OR says at least one must show 1. A gate follows its rule whenever its inputs change.')
picture('gates','Same input cards, different rules: AND gives 0 for 1 and 0; OR gives 1.')
table([['Input B','Input C','AND output','OR output'],[0,0,'____','____'],[0,1,'____','____'],[1,0,0,1],[1,1,'____','____']],[1.35,1.35,2.15,2.15])
tryit('Take turns being an input and the gate. Test all four rows. Add a NOT rule: it flips 0 to 1 and 1 to 0. Draw two input wires going into a box and one output wire leaving it.')
ask('How could you describe the rule so another team could build the same circuit?')
boldline('New idea','An HDL is a hardware description language: a way to describe a circuit and its connections.')
note('Use card play for this lesson or a teacher-prepared gate image. The included Go Board starter implements the adder/counter; it is not a selectable gate demonstration. The answer key supplies the complete truth table.')

# 5 Lesson 3
page('Lesson three  Make an adder')
mission('Connect the statement A = B + 2 to a real circuit.')
picture('adder','When B is 3, the four-bit input is 0011. Add 0010 and the answer is 0101, or 5.')
table([['Input B','Rule','My predicted A','Observed A'],[0,'B + 2','____','____'],[3,'B + 2',5,'____'],[5,'B + 2','____','____'],[7,'B + 2','____','____']],[1.05,1.6,2.3,2.05])
h('Your tiny code edit')
p('In the supplied StudentLogic.cs file, this line builds the adder:')
para=p('var a = b + Expr.Const(b.Width, 2);')
for run in para.runs:run.font.name='DejaVu Sans Mono';run.font.size=Pt(11.5)
para.paragraph_format.space_after=Pt(14)
tryit('Change the final 2 to 3. Predict what happens when B is 3, then rebuild with your guide and test it. The laptop compiles the description; the FPGA becomes the circuit you described.')
ask('Does changing the input B require compiling the circuit again? Does changing the rule require a new build?')
note('For a change to add three, the guide uses boards/GoBoard/lesson.sh 3. This checks against the predicted addend, then generates a fresh image. The source variable a is a reference to an expression, not a hardware register. Drive creates a combinational connection.')

lablink('addtwo','Add two and wrap — eight-bit version')

# 6 Lesson 4
page('Lesson four  What happens when bits run out')
mission('Discover overflow and learn why a fifth bit can matter.')
p('Four bits can represent 0 through 15. The number 16 needs five bits: 10000. Our four-bit answer keeps only the four places on the right.')
picture('overflow','14 + 2 needs five bits. Keeping only the low four bits gives 0000, or zero.')
table([['B','B + 2 in ordinary maths','Four-bit answer A'],[12,14,14],[13,15,15],[14,16,0],[15,17,'____']],[1.0,3.0,3.0])
tryit('Use button 2 to reach B = 14. Predict the right-hand digit. Then reach 15 and press again: the input itself wraps to zero. Use the 8, 4, 2, 1 cards to explain why.')
ask('If we added a place worth 16, which answer could we keep? Where would a fifth LED connect?')
boldline('Display clue','After 9, a hexadecimal display uses A, B, C, D, E and F for ten through fifteen. Use the display key near the end of the workbook.')
note('The raw Expr examples intentionally use modulo arithmetic at their chosen width. Typed UInt addition can preserve a carry in an extra bit; introduce that later, after the child understands the four-bit experiment.')

# 7 Lesson 5
page('Lesson five  A circuit that remembers')
mission('Tell the difference between a changing wire and a stored number.')
picture('memory','The input has changed to 8, but the register holds 3 until its next capture event.',width=6.0)
table([['What the group does','Input card','Stored card'],['Start',0,0],['Change input; no clap',3,0],['Clap',3,3],['Change input; no clap',8,'____'],['Clap',8,'____']],[3.7,1.65,1.65])
picture('clock','A register captures at a rising edge: the moment the clock changes from low to high.',width=6.3)
tryit('One partner holds the input card. The other keeps the stored card and copies the input only when the guide claps. Change inputs between claps. Then try the board’s button counter: a long press should count once.')
ask('Why can an adder’s answer change when B changes, but a stored count wait for an event?')
boldline('New word','A register stores bits. The clock tells it when to capture its next value.')
note('Clap means a rising clock edge. The teacher shell handles synchronization and switch bounce. Press detection gives one increment for each new debounced press.')

# 8 Lesson 6
page('Lesson six  A blinking counter')
mission('Use a slow counting tick to make time visible.')
picture('counter','Count increases on enabled counting ticks. The low bit goes OFF, ON, OFF, ON.')
table([['Count','Four low bits','Rightmost LED'],[0,'0000','OFF'],[1,'0001','ON'],[2,'0010','____'],[3,'0011','____'],[4,'____','____']],[1.0,3.0,3.0])
tryit('Choose counter mode with button 4 and press button 3 to Run. Predict the next count before it changes. Pause it. Clear it with button 1. Clear does not change Run/Pause, so pause first if you want zero to stay visible.')
ask('The low LED changes once per second while counting. Why does one complete ON-and-OFF cycle take two seconds?')
boldline('Clock clue','The board clock is fast: 25 million edges each second. A counting enable lets the stored count change on a much slower tick.')
note('The divider runs continuously. The first increment after Run may occur in less than one second. The two-digit starter shows hexadecimal and is not a precision stopwatch. No button or divided counter signal is used as a new clock.')

lablink('counter','Count, pause, resume')

# 9 Lesson 7
page('Lesson seven  A hardware assembly line')
mission('Track two number cards while two stages work at once.')
picture('pipeline','Card A starts at 3 and becomes 15. Card B starts at 4 and becomes 18.')
table([['After this clap','Add-two register','Times-three register'],[1,'A: 5','empty'],[2,'B: 6','A: 15'],[3,'C: 7','B: 18'],[4,'empty','C: ____']],[1.25,2.8,2.95])
tryit('Give three children the input, add-two and times-three jobs. Move each card across one stage boundary on each clap. Read the two stage boxes just after each clap. The final card can be collected at the following clap if the receiver is ready.')
ask('Why does another stage delay the first answer without stopping the assembly line from making one answer per clap?')
boldline('Two useful words','Latency is how long one card takes. Throughput is how many finished cards you can produce per unit of time.')
note('One C# Map puts add and multiply before one register stage. Two Map calls put a register between them. The supplied tests measure input-accept to output-accept intervals of one versus two cycles. Source variables alone do not create pipeline registers.')

# 10 Lesson 8
page('Lesson eight  Invent a useful circuit')
mission('Design a small project, then write tests before building it.')
picture('hockey','A shot counter is a capstone idea. The illustration is not a supplied decimal scoreboard.')
table([['Test I will perform','I predict','What I observed'],['Clear the count','0','____________'],['Press once for one shot','1','____________'],['Press for two more shots','____','____________'],['Hold the button down','One new count','____________']],[3.35,1.45,2.2])
tryit('Choose a shot counter, score tally or pattern game. Draw inputs, stored values and outputs. Explain what one press means. Give your test plan to a partner and see whether they get the predicted results.')
ask('Which test could reveal an error even when the board appears to be working? What should happen when your number runs out of bits?')
boldline('My project','_____________________________________________________________')
boldline('My circuit remembers','____________________________________________________')
note('The existing starter provides an adder and run/pause counter. A capstone needs edits or a teacher-provided template. Decimal scores, a pattern game and a precision reaction timer are extensions, not features already implemented in the supplied starter.')

# 11 Extra buffer and languages
page('Bonus  When the receiver needs a rest')
mission('Keep every number card even when the next child stops.')
picture('skid','The two-entry skid buffer has a head slot and a spill slot. Its output card stays put during a wait.',width=6.9)
table([['Receiver','Stored cards','What the sender should do'],['Waiting','A in head; spill empty','One extra card may be accepted'],['Waiting','A in head; B in spill','Wait: both slots are full'],['Ready again','A is delivered; B moves to head','Resume when input Ready returns']],[1.35,2.85,2.8],10)
picture('languages','The actual example links separately compiled F#, C# and Visual Basic pieces.',width=6.5)
p('Different source languages can call the same .NET library. They all help construct one hardware graph. CIL is the common software execution format; SharpHDL records the circuit. The FPGA does not run an F# or VB runtime.')
note('The skid buffer adds one cycle of minimum forward latency, not two mandatory cycles. It has a one-cycle input acceptance bubble when leaving the full state. The complete mixed-language example has three cycles minimum latency and can sustain one beat per cycle when unstalled.')

# 12 Printable cards
page('Printable number cards')
p('Choose ON or OFF for each card. ON adds its value; OFF adds zero. Arrange cards in the order 8, 4, 2, 1. Write the bit pattern in the same order.')
picture('printable_cards','Print this page to make a physical binary game.',width=6.5)
table([['Target','Cards that are ON','Binary pattern'],[3,'________________','____'],[6,'________________','____'],[12,'________________','____'],[15,'________________','____']],[1,3.5,2.5])
ask('Can you make every number from zero to fifteen exactly once? Try the same game with just the 2 and 1 cards.')

# 13 Display key
page('The number display detective key')
picture('segments','A digit is made by lighting selected segments. The letter labels name segments, not bit positions.',width=6.7)
p('The starter uses hexadecimal: one digit for each four-bit value. The picture shows which segments form 2. The real board uses active-low segment pins; the teacher-owned shell handles that inversion.')
rows=[['Decimal','Hex','Binary','Decimal','Hex','Binary']]
for i in range(8):rows.append([i,format(i,'X'),format(i,'04b'),i+8,format(i+8,'X'),format(i+8,'04b')])
table(rows,[1.05,.7,1.75,1.05,.7,1.75],11)
tryit('Find binary 1011. Which hex digit represents it? Find F. What decimal number is it? Why can one hex digit always fit exactly four bits?')

# 14 Teacher plan
page('Guide for the grown up builder')
p('Start with one board per pair and one source language. Prepare the shell, editor and USB programming before class. Keep the first edit to a constant in StudentLogic.cs.')
table([['Minutes','Activity','Guide’s job'],[10,'Predict the result','Ask for a reason before the answer'],[10,'Cards or a physical model','Make bits and registers tangible'],[25,'Build and try','Handle tool setup and programming'],[10,'Test predictions','Keep automatic checks in the build'],[5,'Explain and reflect','Have partners swap explanations']],[.8,2.7,3.5])
h('The supported classroom target')
p('Use .NET 8 SDK, Python, Icarus, Yosys, nextpnr-ice40 and IceStorm. OSS CAD Suite packages the FPGA tools. The supplied shell scripts require a POSIX shell; prepare the environment and USB drivers for your computers. The small student file calls the same explicit-width expression API as the larger examples.')
p('The teacher-owned board shell handles startup reset, two-register input synchronization, 10 ms debouncing, pin constraints and segment encoding. The C# application logic supplies B, the adder, a counter and display mode. Use clock enables; do not create a clock from a button.')
h('What was actually checked')
p('The existing examples passed 3,002 .NET assertions, seven generated-Verilog benches and a board-shell simulation. The classroom target completed placement/routing at 25 MHz and bitstream packing: 230 of 1,280 logic cells, no block RAM. No physical board has been connected or tested. The larger floating-point library is outside this small-board course.')
h('Before equipping the group')
p('Validate startup, button polarity, single increments on holds, both displays, Run/Pause and a changed-constant rebuild on one real board. Confirm its USB programming interface and revision. For twelve children, six stations budget about $540 before tax/shipping: six $70 boards plus estimated $20 accessories per station. Computers are excluded.')
h('Sources and companion files')
for text in [
 'Board price/features checked 13 October 2026: nandland.com/product/the-go-board/ and nandland.com/the-go-board/',
 'Free FPGA tool distribution: github.com/YosysHQ/oss-cad-suite-build',
 'Place and route / bitstream tools: github.com/YosysHQ/nextpnr and github.com/YosysHQ/icestorm',
 'Companion: SharpHDL 0.13, boards/GoBoard and examples/ThreeLanguages. Original diagram assets and builders are included; cover art is generated illustration.'
]:
 para=p(text)
 for run in para.runs:run.font.size=Pt(9)

# 15 Answers
page('Answer key and progress passport')
p('Use the answers after children have made predictions. Praise a clear explanation and a test that catches a mistake, rather than the number of lines typed.')
table([['Activity','Answer'],['Lesson 1','6 = 0110. A bit has 2 choices. Four bits have 16 patterns.'],['Lesson 2 AND','00 -> 0; 01 -> 0; 10 -> 0; 11 -> 1.'],['Lesson 2 OR','00 -> 0; 01 -> 1; 10 -> 1; 11 -> 1.'],['Lesson 3 add two','B = 0, 3, 5, 7 gives A = 2, 5, 7, 9.'],['Lesson 4','B = 15 gives A = 1 after adding two in four bits.'],['Lesson 5','Change to 8 without a clap: hold 3. Next clap: store 8.'],['Lesson 6','Count 2: OFF; count 3: ON; count 4: 0100 and OFF.'],['Lesson 7','C = 5 -> 7 -> 21. One boundary is crossed per clap.'],['Lesson 8','After one shot, two more give 3. A new held press adds one count.'],['Number cards','3: 2+1 = 0011; 6: 4+2 = 0110; 12: 8+4 = 1100; 15: all = 1111.'],['Display detective','1011 is hex B and decimal 11. F is decimal 15.']],[1.65,5.35],10.5)
h('My hardware passport')
table([['I can show or explain','Date','Partner’s initials'],['Read a four-bit number','________','________'],['Predict AND and OR','________','________'],['Explain a four-bit overflow','________','________'],['Explain why a register remembers','________','________'],['Predict the blinking low bit','________','________'],['Track two pipeline cards','________','________'],['Write a test that catches a mistake','________','________']],[4.3,1.35,1.35],10.5)
p('For new typists, the first six lessons can fill the course. Spend longer on cards and physical models. CIL, generics, floating point and Verilog syntax are optional later topics, not entry requirements.')

# 16 Teacher operator reference
page('Teacher reference  Supported operators')
p('This table describes the explicit Expr library. A separate CilCompiler frontend now translates selected ordinary int, uint and bool methods, including comparisons and branches. The teacher extensions explain that second route.')
table([['Raw Expr operation','Hardware meaning and width'],
 ['a + b, a - b, a * b','Add, subtract, multiply. Equal W-bit operands give W bits; overflow wraps. Widen inputs first for carry or full product.'],
 ['a / b, a % b','Unsigned quotient and remainder, W bits. Toward zero; division widths 1 to 256. Explicit API exposes error flags.'],
 ['a & b, a | b, a ^ b, ~a','Bitwise AND, OR, XOR, NOT. Width stays W.'],
 ['a << n, a >> n','Constant shifts; right shift fills with zero. Width stays W.'],
 ['a.ArithmeticShift(n)','Constant right shift that repeats the sign bit.'],
 ['a.Eq(b), a.Lt(b), a.SignedLt(b)','Equality, unsigned less-than, signed less-than. Return one bit.'],
 ['Expr.Mux(test, yes, no)','Hardware selector. Test is one bit; both choices have equal width.'],
 ['Resize, SignExtend','Resize keeps low bits or zero-extends. SignExtend widens signed values.'],
 ['Slice, Expr.Concat','Select a bit range or join high bits to low bits.'],
 ['Expr.Lookup','Constant lookup table. An out-of-range address returns zero.'],
 ['Module.Next, NextWhen','Explicit clocked storage; NextWhen adds an enable.']],[2.65,4.35],10)
h('Types can preserve the extra bits')
p('UInt<W> and SInt<W> support +, -, *, / and %. Add/subtract widen to W+1; multiply widens to 2W; division keeps W. Unsigned subtraction returns a signed result. SFixed<W,F> also tracks fractional bits: a product has 2W total and 2F fractional bits. Resize or Convert chooses wrapping or saturation; fixed conversion also chooses rounding. SFixed << wraps at the same width; >> rounds down. ShiftRight can select NearestEven.')
h('Boundaries to teach explicitly')
p('Raw Expr uses constant shifts and Mux for choices. CilCompiler supports variable integer shifts and ordinary if; runtime loops use explicit BuildSequential. Direct fixed-point division is not implemented. Typed integer wrappers expose their Bits for logic and comparisons. The separate compact-float library has arithmetic and math functions; it is not IEEE floating point and is outside this starter course.')
note('The complete reference is OPERATORS_AND_MULTIPLIERS.md. CORDIC, ROM math, streams and buffering are library blocks. Their latency is part of the interface; source variables and comments do not insert registers.')

# 17 Multiply choice and activity
page('Teacher extension  Two ways to multiply')
p('Integer multiply is implemented. The new architecture option chooses a parallel circuit or a circuit that reuses an arithmetic datapath over several clocks. Both take one packed operand pair and produce a full-width product.')
table([['Implementation','Parallel','BitSerial'],['Input','One pair {B,A}','One pair {B,A}'],['Work','Combinational multiply then register','One multiplier bit each compute clock'],['8-bit operands','16-bit product','16-bit product'],['Earliest output acceptance','1 clock after input acceptance','9 clocks after input acceptance'],['Best input interval','Every clock','Every 9 clocks for 8-bit operands'],['While receiver waits','Hold output; stop when occupied','Hold output; stop when occupied']],[2.2,2.4,2.4],10)
p('Select the architecture with an explicit library option. This is a pragma-like instruction to the circuit builder, shared by C#, F# and Visual Basic. Custom C# #pragma text is not interpreted by SharpHDL; merely adding an attribute or a comment would not change the hardware.')
code('var product = IntegerMultiply.Build(\n    m, pairs, "multiply", 8,\n    MultiplyArchitecture.BitSerial);')
note('pairs is a 16-bit stream: A is bits 7:0 and B is bits 15:8. Use signed: true for two’s-complement operands. TKEEP must be all ones. TLAST and TUSER travel with the result. Widths supported here are multiples of four from 4 to 2048; practical FPGA limits are smaller.')
h('A four clap multiplication game')
p('Try 3 times 5 with four-bit numbers. Write B = 0101 and read its bits from right to left. Add the shifted A only when that B bit is 1. Loading the cards takes a separate edge.')
table([['Compute clap','B bit','Shifted A','Total after clap'],[1,1,3,3],[2,0,6,3],[3,1,12,15],[4,0,24,15]],[1.7,1.1,1.7,2.5],10)
note('The result becomes valid after the fourth compute edge and can be accepted on the next edge. Parallel does not guarantee a DSP block; synthesis chooses mapping. Serial is not always smaller at tiny widths because it also needs registers and control. These multipliers have passed simulation and RTL checks, not a physical-board test.')

# 18 Division and rounding
page('Teacher extension  Divide and round')
p('Division can also be a circuit. A parallel divider lays out the comparison and subtraction steps together. An iterative divider reuses a step for one quotient bit per clock. Both now support signed and unsigned integers.')
table([['Rounding choice','7 divided by 3','Negative 7 divided by 3'],['TowardZero','2','-2'],['Floor','2','-3'],['Ceiling','3','-2'],['NearestEven','2','-2'],['AwayFromZero','3','-3']],[2.3,2.3,2.4],10)
p('NearestEven picks the nearest whole number. Halfway cases choose an even answer: 5 divided by 2 rounds to 2, while 7 divided by 2 rounds to 4. Exact divisions do not change with the rounding choice.')
h('Keep the leftover consistent')
p('The residual is A minus Q times B. For 7 divided by 3, choosing Q = 2 leaves 1. Choosing Q = 3 leaves -2. We use an extra signed bit for this leftover, even when the input numbers are unsigned.')
code('var divided = IntegerDivide.Build(\n    m, pairs, "divide", 8,\n    DivideArchitecture.Iterative,\n    DivisionRounding.NearestEven,\n    signed: true, overflow: Overflow.Saturate);')
p('Here pairs is a 16-bit input stream, packing two eight-bit numbers. The output is 24 bits: quotient in bits 7:0, signed residual in 16:8, divide-by-zero in bit 17, overflow in bit 18, and zero padding in 23:19. TLAST and TUSER follow the answer.')
table([['Question','Defined behavior'],['When is the answer accepted?','Parallel: next edge. Eight-bit iterative: nine edges after input acceptance, if the receiver is ready.'],['What if B is zero?','Quotient 0, residual A, divide-by-zero flag set.'],['What if signed MIN is divided by -1?','Overflow flag set. Wrap gives MIN; Saturate gives MAX. Residual is 0.'],['What do raw Expr / and % do?','Combinational, toward zero, Wrap, with no status output. Use the explicit API when flags matter.']],[2.3,4.7],10)
note('Iterative division computes for eight clocks after loading; it holds the result while the receiver waits. Reset discards pending work. These are explicit hardware conventions, not CPU divide exceptions. Direct SFixed division and selectable compact-float rounding are not added by this feature.')

# 19 Rationale and synthesis limits
page('Why SharpHDL')
p('SharpHDL lets us use familiar programming tools to build real circuits. C#, F# and Visual Basic can share the same library. Functions and construction loops can create repeated channels, packet fields, arithmetic and stream connections without copying whole designs by hand.')
p('For children, the benefit is a small source file, visible numbers and an explicit model of wires and memory. For engineers, the benefit is reusable circuit construction, simulation and readable Verilog. Hardware still needs clocking, widths, latency, verification and an implementation that fits the board.')
h('What the software compiler contributes')
p('The language compiler checks names and types and creates CIL. SharpHDL either runs an explicit builder or analyzes a selected scalar method. Its new CIL analyzer translates comparisons and branches into logic and selectors. Both routes emit Verilog for FPGA tools. The FPGA does not run .NET; pipelines still need explicit boundaries.')
h("What it won't synthesize")
table([['Request','Current boundary'],['Any ordinary .NET program','CilCompiler accepts a restricted typed subset. General objects, double and arbitrary programs remain unsupported.'],['An unbounded hardware loop','Ordinary if/switch and constant loops now translate. BuildSequential compiles supported runtime loops with a watchdog.'],['Objects, strings, tasks or file I/O on the FPGA','Host facilities may help build a circuit; no .NET runtime or filesystem is generated.'],['Automatic math replacement','Call the explicit CORDIC, ROM or compact-float block. System.Math is not automatically replaced.'],['Automatic pipeline or shared arithmetic','Insert registers or select an iterative block explicitly. Source variables do not mark clock boundaries.'],['Full IEEE floats or direct fixed-point divide','The compact format has its own rules; SFixed division is not implemented.'],['A complete FPGA system from any graph','Single-port RAM is now explicit. Other memory IP, clock domains, pins, constraints and timing need additional design.']],[2.25,4.75],10)
p('This is an executable teaching and prototyping library. Supported logic can still be too large or too slow for a particular FPGA. The Go Board starter has completed software placement and routing; physical-board validation remains to be done.')
note('The technical reference SYNTHESIS_BOUNDARIES.md expands these limits. INTEGER_DIVISION.md specifies rounding, result buses, error flags and timing. The workbook keeps these topics as optional teacher material, not prerequisites for the first lessons.')

# 20 Behavioral extension
page('Teacher extension  Ordinary code becomes gates')
p('An ordinary comparison can now become hardware. Start with two number cards, A and B. Ask a child to choose the smaller one. A comparator makes the decision, and a selector sends the chosen card to the output.')
code('public static int Minimum(int a, int b) {\n    if (a < b) return a;\n    return b;\n}')
table([['Step','What happens'],['C# compiler','Checks types and emits loads, a conditional branch and returns.'],['SharpHDL CIL analyzer','Loads symbolic wires, compares signed values and explores both paths.'],['Hardware graph','A comparison controls a selector. No register is implied.'],['FPGA tools','Map the Verilog to the selected chip’s logic resources.']],[1.8,5.2],10.5)
h('Three languages share the middle step')
p('F# can write if a < b then a else b. Visual Basic can write If a < b Then Return a followed by Return b. Their compiled methods all use CIL, so the same hardware analyzer can handle them. Different overflow rules still matter: checked arithmetic produces a fault signal.')
h('Wire functions into a stream')
p('The tested example takes ordinary F# Minimum, wraps it in a clocked stream Map, inserts the C# skid buffer, and adds a VB arithmetic Map. Result bits 31:0 and fault bit 32 travel in a 40-bit data bus. Bits 39:33 are zero. Packet markers travel alongside the answer.')
table([['Question','Answer'],['Does a method call add a clock?','No. The explicit Map or register creates the boundary.'],['What happens when the receiver pauses?','The stream and skid buffer retain pending values and metadata.'],['What did the tests compare?','Actual .NET method answers and faults against the circuit simulator and emitted RTL.']],[2.8,4.2],10.5)
note('Teacher reference: BEHAVIORAL_CIL.md and examples/Behavioral. The frontend is a tested restricted prototype. BuildSequential adds explicit operation-by-operation clocking; arbitrary .NET programs remain unsupported.')

# 21 Loop and memory extension
page('Teacher extension  Loops shifts and memory')
p('A loop can mean copies of hardware or work spread across clock ticks. Decide which meaning you want before choosing how many bits the counter needs.')
table([['Loop style','Where did the counter go?'],['Repeat twelve times while building a circuit','The compiler unrolls twelve copies. There is no hardware loop counter.'],['Count clock ticks from 1 through 12','Keep an explicit register. Four unsigned bits fit 12 and the terminal value 13.'],['Use a signed counter','Four signed bits only reach +7. Positive 12 needs five signed bits.']],[2.9,4.1],10.5)
p('SharpHDL removes the constant index in its tested twelve-iteration example. It does not yet infer the smallest safe width for clocked state. A future range check must include reset, increment and exit values before suggesting a narrower counter.')
h('A shift is a number slide')
table([['Operation','Example','Hardware rule'],['Multiply by two','0011 becomes 0110','Shift left; check overflow if requested.'],['Unsigned divide by two','0111 becomes 0011','Shift right, dropping the last bit.'],['Signed divide by two','-3 divided by 2 is -1','Correct negative rounding before the shift.']],[2.1,2.05,2.85],10.5)
h('RAM is a row of numbered drawers')
p('An address chooses a drawer. Each drawer holds one word. Our new RAM has read and write enables and one shared address. On a clock clap, a read copies a drawer into the output register. With read disabled, the old output stays visible.')
table([['At the clock clap','Output after the clap','Drawer after the clap'],['Read a drawer holding 7','7','7'],['Read and write 9 together','7 — the old value','9'],['Reset','0','9 — contents retained']],[2.5,2.25,2.25],10.5)
note('Try it with envelopes and number cards. The simultaneous read/write rule is read-before-write. MEMORY.md specifies initialization, address bounds and reset. A 256-by-16 example mapped to one FPGA block RAM and passed mapped simulation; it has not been tested on a physical board.')

# 22 Interactive companion
page('Clock lab  Coach the circuit')
p('Open SharpHDL_Learning_Lab.html in a modern browser. It works offline. This companion is most useful after Lesson 5, when children know that a register changes on a clock edge. Start with the hockey drill; the number shifts and loops are optional extensions.')
picture('fsm_lab','The controller remembers one state. Start enters Running; four enabled practice ticks reach Done. Stop cancels a running drill. Reset returns to Ready.',width=6.2)
h('Predict before you press the clock')
table([['Try this','Predict and explain'],['Turn Start on and press Clock once','Running begins, but count is still zero. The first practice tick is the next edge.'],['Turn Start off and clock four times','Count reaches four and state becomes Done. One more edge returns to Ready.'],['Pause by turning Practice enabled off','Running and count hold until practice is enabled again.'],['Turn Stop on during a drill','The next edge returns to Ready. Reset also clears the count immediately on its edge.']],[2.9,4.1],10.5)
h('A coach with a very small memory')
p('Ready, Running and Done fit in a two-bit state register. The helper generates its next-state choices. The count is a separate four-bit register. Both update together from their old values: writing one does not instantly change the other within the same clock tick.')
code('var fsm = new StateMachine(m, "drill",\n    "Ready", "Running", "Done");\nfsm.When("Ready", start, "Running")\n   .When("Running", finished, "Done");\nfsm.Build();')
p('This short extract shows the idea; the HTML includes the full C# with Stop and Done transitions. A missing transition holds the state. When two transitions match, the first listed wins. The helper creates real hardware; it does not automatically schedule an arbitrary while loop.')
note('The Clock Lab simulates exported graphs. The linked WASM lab runs generated Verilog with a teacher bench. Both are checked against expected traces. Fresh WASM runs need engine files.')

lablink('fsm','Coach the hockey drill in real RTL')

# 23 Sequential extension
page('Teacher extension  A loop with a clock')
p('Imagine a robot following instruction cards. One clap lets it follow one tiny instruction. A loop sends it back to an earlier card. It keeps its numbers in registers between claps.')
code('public static int Sum(int n, int total) {\n    while (n > 0) { total += n; n--; }\n    return total;\n}\n// Select CilCompiler.BuildSequential for clocked work.')
table([['BuildModule: unroll','BuildSequential: clocked'],['Known repetitions become a combinational graph.','Runtime conditions choose the next instruction state.'],['No clock for each iteration.','One CIL operation per clock, not one whole iteration.'],['Repeated work may make a long path.','Registers remember work between edges.']],[3.5,3.5],10.5)
h('Ask the robot to add 3 + 2 + 1')
table([['Signal card','Child-friendly meaning'],['Start','Please begin! Inputs are copied when the robot is idle.'],['Busy','I am working. Another Start is ignored.'],['Done','Finished! This stays on until reset or another accepted Start.'],['Fault / timeout','Something went wrong, or the instruction budget ran out.']],[1.6,5.4],10.5)
p('Try seed 10 and n = 3. Predict 16, then open A loop that takes clock ticks in the workbook WASM lab. Run the prepared bench and inspect busy, done and result. The adult owns the testbench; children predict and read the waveform.')
note('Timing depends on compiled CIL: loads, comparisons, branches and stores take separate clocks. This first scheduler does not promise the smallest circuit, automatic pipelining or a particular FPGA frequency. A memory read adds one wait edge.')

lablink('sequential','A loop that takes clock ticks')

# 24 Typed values and safety
page('Teacher extension  Number boxes and drawers')
p('A twelve-bit unsigned number box has room for 0 through 4095. Add one to 4095 and it wraps to zero. The UBits<W12> type makes that happen at each operator, in both ordinary software execution and generated hardware.')
code('UBits<W12> sum = a + b;\nreturn sum >> 1;  // First wrap the sum, then halve it.')
table([['Concept','Try it with paper'],['Typed width','Use four cards to make 1111. Add one: only four cards fit, so the result is 0000.'],['Packet fields','Tape twelve length squares, eight kind squares and one urgent square into a 21-bit row. Field offsets tell us where to look.'],['RAM array','Number sixteen envelopes 0 to 15. A uint[] parameter in sequential mode represents these fixed drawers.'],['Assertion','Before opening a drawer, check the rule. Hardware.Assert turns a failed rule into a visible fault.']],[1.55,5.45],10.5)
h('Important rules for the adult')
p('The packet example reads a flat, explicitly laid-out header. The RAM example owns one uint[] memory with a fixed depth. General objects, nested packet records and arbitrary arrays remain outside this first subset.')
p('Reset stops the robot and clears its control registers, but the memory drawers keep their contents. A fault does not undo earlier writes. Done is a held signal, not a one-clock flash. Read result only when Done is on and Fault is off.')
h('Two complementary experiments')
p('The offline Clock Lab evaluates exported circuit graphs. The separate workbook WASM lab runs actual generated Verilog and prepared testbenches on the page, using the VeriSim Icarus WebAssembly engine. Its five experiments have browser and native-trace agreement tests. Fresh runs need engine files; bundled recorded traces work offline.')
note('Requirements and exact boundaries: CONSTRUCTS_AND_REQUIREMENTS.md and SEQUENTIAL_CIL.md. Reproducible native-code, graph and RTL checks: tests/run-advanced.sh. These tests establish the included cases; they are not a proof for every program.')

# 25 Traceable generated hardware
page('Teacher extension  Names on the wires')
p('A label helps us follow a wire through a drawing. SharpHDL can put useful labels on the wires in its generated Verilog too. The label changes what we call the wire, not the number it carries or when that number changes.')
code('''var address = (baseAddress + offset)
    .Named("effective_address");
var selected = address.Eq(Expr.Const(16, 0xff02))
    .Named("uart_tx_selected");''')
h('Follow a labeled wire')
code("""assign n_uart_tx_selected =
    n_effective_address == 16'd65282;""")
p('Read this aloud: the UART transmit register is selected when the address equals 65282. The same number written in hexadecimal is FF02. In TinyRisc, that address means send a byte. The comparison makes a one-bit yes-or-no wire; it does not store a number or add a clock.')
table([['Label','What it tells us'],['n_effective_address','The address calculated by adding the base and offset.'],['n_uart_tx_selected','A yes-or-no result identifying the transmit register.'],['n_add_17','An unnamed addition; the number distinguishes this particular wire.'],['Method_IL_0002_add','Arithmetic traced to a method and instruction offset in CIL.']],[2.5,4.5],10.5)
h('Try a detective experiment')
p('Open docs/tinyrisc/TinyRisc_Demo.html. Step one instruction at a time and watch the PC, registers and output LEDs. Predict the next value before stepping. The prepared program adds 5 + 4 + 3 + 2 + 1 and leaves 15 on the LEDs. Its bundled trace works offline; the optional fresh VeriSim run needs the simulator engine.')
note('The source builder supplies application names explicitly with Named. C# local names are not automatically recovered. Duplicate hints receive suffixes; every intermediate wire keeps its width and truncation. The demo steps through a recorded trace. A fresh VeriSim run executes real RTL, but is not a live paused processor. Requirements and examples: NET_NAMING.md.')

# 26 Beginner waveform experiments
page('Browser lab  Follow the clock')
p('Open the workbook WASM lab after your guide starts the server described on the next page. WASM means WebAssembly: it lets a simulator run in the browser. You do not need to install FPGA software to explore the recorded traces.')
h('Predict first and run second')
p('Choose Add two and wrap. Predict the answers for B = 0, 7 and 254. Press Run RTL with WebAssembly. When it reports a fresh simulation passed, move the cursor and read A. The answers are 2, 9 and 0. This browser example uses eight bits; our starter-board arithmetic lesson uses four.')
lablink('addtwo','Add two and wrap')
h('Now give a register a clock')
p('Choose Count, pause, resume. Look for the upward edges of clk. Enable lets the counter move; a low enable makes it keep its number. The orange line is your viewing cursor, not another signal.')
picture('wasm_counter_wave','Actual browser simulation: reset to zero, count three, hold for two clocks, then count once more.',width=6.8)
table([['Rising edge time','What happens','Predict the count'],['2 ns','Reset','0'],['6, 10 and 14 ns','Enable is on','1, 2, 3'],['18 and 22 ns','Enable is off','____, ____'],['26 ns','Enable is on again','____']],[1.7,2.9,2.4],10.5)
lablink('counter','Count, pause, resume')
ask('Why does the adder change while clk stays low, but the counter wait for an upward edge?')
note('Counter answers: 3, 3, 4. Next event may visit a falling edge or an input change; it is not Next clock. First/Next/Last event replay a completed trace. The time labels arrange the test; they are not measurements of FPGA speed.')

# 27 Teacher setup and extensions
page('Browser lab  Guide and extension cards')
h('Start the lab on a computer')
code('python3 -m http.server 8090 --directory docs')
p('Run that command from the repository root. Open http://localhost:8090/interactive/SharpHDL_WASM_Lab.html on the same computer. Workbook experiment links use this address. For a tablet, replace localhost with the teacher computer’s address.')
p('The default fresh simulation loads the VeriSim engine from the internet. To prepare offline use, first run the command below while connected, then check Use downloaded engine in the lab. Keep the supplied docs folders together.')
code('python3 tools/fetch_verisim_runtime.py')
h('Three extension cards')
table([['Experiment','Predict before running','Look at these waves'],['Coach the hockey drill','Start begins at zero; four following enabled edges reach Done.','s_ticks, s_phase, s_done'],['Twelve additions, no clocks','2 + 12 × 3 = ____. Does it need clock edges?','s_a, s_b, s_y; clk stays low'],['A loop that takes clock ticks','10 + 3 + 2 + 1 = ____. When is it valid?','s_busy, s_done, s_result, s_fault']],[2.0,3.0,2.0],10.2)
lablink('fsm','Coach the hockey drill')
lablink('forloop','Twelve additions, no clocks')
lablink('sequential','A loop that takes clock ticks')
h('What the buttons really do')
p('Run RTL compiles and simulates the generated Verilog in WebAssembly, then replaces the waveform trace. The teacher’s testbench checks the answers. Restore recorded trace returns to the bundled native-Icarus result. Cancel stops a worker; a run also has a 60-second limit. The children can read waves without writing a testbench.')
note('Extension answers: four enabled ticks; 38 without clock edges; 16 when Done is high and Fault and Timeout are low. Unknown X/Z is not zero. A CIL operation can take a clock; a whole loop iteration usually takes several. Full source, testbenches, requirements and engine attribution: WASM_WORKBOOK.md. TinyRisc_Demo.html is the next PC/register/LED exercise.')

# Remove inherited title rules and keep all headings black.
for tree in (d.styles.element,d.element):
 for node in list(tree.iter(qn('w:pBdr'))):node.getparent().remove(node)
output=Path(sys.argv[1] if len(sys.argv)>1 else 'SharpHDL_Classroom_Proposal.docx');d.save(output)
print(output.resolve())
